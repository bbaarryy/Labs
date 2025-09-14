import matplotlib.pyplot as plt
import seaborn as sns

from docx import Document

doc = Document()

table = doc.add_table(rows=3, cols=4)
table.cell(0, 0).text = "Ошибка"
table.cell(1, 0).text = "g1"
table.cell(2, 0).text = "g1"
table.cell(0, 1).text = "Число случаев"
table.cell(0, 2).text = "Доля случаев"
table.cell(0, 3).text = "Теоретическая оценка"

table.style = 'Light Shading Accent 1'

doc.save('demo.docx')

# arr80 = get_arr(raw,20)
# mx = max(arr80)

# all = 0

# x = [1]
# for i in range(2,mx+1):
#     x.append(i)
# y = [0] * mx

# for i in range(len(arr80)):
#     if(arr80[i] != 0):
#         y[arr80[i]-1]+=1
#         all+=1

# for i in range(len(y)):
#     y[i] = y[i]/all

# ax = plt.gca()
# #ax.axes.xaxis.set_visible(False)
# vis = []
# # e = [""]
# for i in range(0,mx,10):
#     for j in range(9):
#         vis.append("")
#     vis.append(str(i+10))
# ax.set_xticklabels(vis)
# sns.barplot(x=x, y=y, color='blue')
# plt.xlabel("n") # Подпись оси X
# plt.ylabel("w") # Подпись оси Y
# plt.title('Гистограмма для t = 20') #Название


# plt.show()